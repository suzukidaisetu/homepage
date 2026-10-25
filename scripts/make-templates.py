"""Makes the reference sheet (.xlsx) and the texts (.docx) from a built site.json, for moving to the sheet/doc format.

    python3 scripts/make-templates.py <site.json> <out dir> [--files content/references]

Upload the results to the Drive folder: リファレンス.xlsx to the top of homepage/, the .docx files into texts/.
(Open each in Google Sheets / Docs and use File > Save as Google Sheets / Docs, or turn on
"Convert uploads" in Drive settings. The site reads both the converted files and the original .xlsx/.docx.)
"""
import json, re, sys, os, html
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from docx import Document
from docx.shared import Pt

site = json.load(open(sys.argv[1]))
out = sys.argv[2]
os.makedirs(out, exist_ok=True)
files = []
if '--files' in sys.argv:
    d = sys.argv[sys.argv.index('--files') + 1]
    files = sorted(f for f in os.listdir(d) if re.search(r'\.(jpe?g|png|webp|heic|pdf)$', f, re.I))

works, refs, texts = site['works'], site['refs'], site['texts']
plain = lambda s: html.unescape(re.sub(r'<br\s*/?>', '\n', re.sub(r'<(?!br)[^>]+>', '', s or '')))
work_name = lambda w: f"{w['no']}_{plain(w['title'])}"

# words that link a reference inside the texts by themselves (kept from the old [word](04) / refs: links)
WORDS = {'sakimori — Koukei': '後景, Koukei'}

wb = Workbook()
ws = wb.active
ws.title = 'リファレンス'
head = ['見た目', '画像', '題名', '種類', 'タグ', '作品', '言葉', '一言', '一言(英)', 'リンク', 'ID']
widths = [16, 26, 28, 12, 26, 34, 18, 44, 44, 40, 14]
ws.append(head)
for k, r in enumerate(refs):
    n = int(r['no'][4:])
    fname = files[k] if k < len(files) else f"{n:02d}_{plain(r['title'])}.jpg"
    note = r.get('note') or {}
    ws.append(['', fname, plain(r['title']), plain(r['kind']), ', '.join(plain(t) for t in r['tags'] if t != 'Other'),
               ', '.join(work_name(works[i]) for i in r['rel']), WORDS.get(plain(r['title']), ''),
               '\n\n'.join(plain(p) for p in note.get('ja', [])), '\n\n'.join(plain(p) for p in note.get('en', [])),
               '\n'.join(l['url'] for l in r['links']), ''])
for i, w in enumerate(widths):
    ws.column_dimensions[chr(65 + i)].width = w
ws.freeze_panes = 'C2'
for c in ws[1]:
    c.font = Font(bold=True)
    c.fill = PatternFill('solid', fgColor='EEEEEA')
for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
    for c in row:
        c.alignment = Alignment(vertical='top', wrap_text=True)

lists = {
    '作品一覧': [work_name(w) for w in works],
    'タグ一覧': sorted({plain(t) for r in refs for t in r['tags'] if t != 'Other'} | {'Architecture', 'Drawing', 'Model', 'Photo', 'Music', 'Book', 'Film', 'City', 'Landscape', 'Art'}),
    '種類一覧': sorted({plain(r['kind']) for r in refs} | {'DRAWING', 'MODEL', 'PHOTO', 'MUSIC', 'BOOK', 'FILM', 'EXHIBITION', 'PLACE', 'IMAGE'}),
}
for name, vals in lists.items():
    s = wb.create_sheet(name)
    s.append([name.replace('一覧', '')])
    s['A1'].font = Font(bold=True)
    for v in vals:
        s.append([v])
    s.column_dimensions['A'].width = 40
for col, src in [('D', '種類一覧'), ('E', 'タグ一覧'), ('F', '作品一覧')]:
    dv = DataValidation(type='list', formula1=f"'{src}'!$A$2:$A$300", allow_blank=True, showErrorMessage=False)
    ws.add_data_validation(dv)
    dv.add(f'{col}2:{col}400')

how = wb.create_sheet('使い方')
for line in [
    'リファレンス表の使い方',
    '',
    '1行が1つのリファレンスです。行の順番がサイトの並び順になります(行はドラッグで入れ替えられます)。',
    '画像:references フォルダに入れた画像のファイル名(拡張子や先頭の番号は省略してもよい)。',
    '題名:サイトに出る名前。空なら画像のファイル名から。',
    '種類・タグ・作品:一覧から選ぶ(複数のときは「, 」で区切る)。作品は works のフォルダ名か番号。',
    '言葉:テキストの本文にこの言葉が出てきたら、自動でこのリファレンスに結び付く(「, 」で複数)。',
    '一言/一言(英):リファレンスの詳しい欄に出る文章。セルの中で改行してよい(Mac: ⌘+Enter、Windows: Ctrl+Enter)。空行で段落。',
    'リンク:Spotify / Apple Music / YouTube などのURL。複数は改行で。',
    '見た目・ID:メニュー「サイト」→「画像と作品を取り込む」で自動で入る(手で書かなくてよい)。',
    '',
    '表に書いていない画像も、サイトにはそのまま出ます。',
]:
    how.append([line])
how.column_dimensions['A'].width = 110
how['A1'].font = Font(bold=True, size=13)
wb.save(os.path.join(out, 'リファレンス.xlsx'))

# ---- texts ----
def para(doc, text, style=None):
    p = doc.add_paragraph(style=style)
    p.add_run(text)
    return p

for t in texts:
    no = t['no'][2:]
    doc = Document()
    doc.styles['Normal'].font.size = Pt(11)
    first = True
    for lang in ['ja', 'en']:
        x = t.get(lang)
        if not x:
            continue
        if not first:
            doc.add_paragraph()
            para(doc, 'English' if lang == 'en' else '日本語', 'Heading 1')
        para(doc, plain(x['title']), 'Title')
        if x.get('kind'):
            para(doc, ('種類: ' if lang == 'ja' else 'kind: ') + plain(x['kind']))
        if first:
            if t.get('date'):
                para(doc, '日付: ' + plain(t['date']))
            if t.get('works'):
                para(doc, '作品: ' + ', '.join(work_name(works[i]) for i in t['works']))
            extra = [plain(refs[j]['title']) for j in t.get('refs', []) if not WORDS.get(plain(refs[j]['title']))]
            if extra:
                para(doc, 'リファレンス: ' + ', '.join(extra))
            if t.get('draft'):
                para(doc, '下書き: yes')
        for b in x['body']:
            para(doc, plain(b))
        first = False
    title = plain((t.get('ja') or t.get('en'))['title'])
    doc.save(os.path.join(out, f'{no}_{title}.docx'))
print('written to', out)
